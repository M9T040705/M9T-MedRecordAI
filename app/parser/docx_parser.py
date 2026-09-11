"""DOCX 解析器：python-docx 抽取段落 + 表格。"""
from __future__ import annotations

from pathlib import Path

from .base import BaseParser, ParsedDocument, register


@register
class DocxParser(BaseParser):
    extensions = ["docx", "doc"]

    def parse(self, path: Path) -> ParsedDocument:
        from docx import Document

        doc = Document(str(path))
        parts, tables = [], []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())
        for table in doc.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            tables.append(rows)
            for row in rows:
                parts.append(" | ".join(c for c in row if c))
        return ParsedDocument(doc_name=path.name, format="docx",
                              text="\n".join(parts), tables=tables)
