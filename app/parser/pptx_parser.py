"""PPTX 解析器：python-pptx 抽取幻灯片文本与表格。"""
from __future__ import annotations

from pathlib import Path

from .base import BaseParser, ParsedDocument, register


@register
class PptxParser(BaseParser):
    extensions = ["pptx", "ppt"]

    def parse(self, path: Path) -> ParsedDocument:
        from pptx import Presentation

        prs = Presentation(str(path))
        parts, tables = [], []
        for idx, slide in enumerate(prs.slides, start=1):
            slide_text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                    if t:
                        slide_text.append(t)
                if shape.has_table:
                    rows = [[c.text.strip() for c in row.cells] for row in shape.table.rows]
                    tables.append(rows)
                    slide_text.append(" | ".join(" | ".join(r) for r in rows))
            if slide_text:
                parts.append(f"[幻灯片{idx}]\n" + "\n".join(slide_text))
        return ParsedDocument(doc_name=path.name, format="pptx",
                              text="\n\n".join(parts), tables=tables)
