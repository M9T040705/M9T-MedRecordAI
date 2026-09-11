"""病案首页样例生成器：生成 txt/md/pdf/docx/xlsx/png 6 种载体 + expected.json。

- 12 份样例记录（字段值覆盖常见组合，ICD 编码均可在 icd_sample.json 命中）；
- expected.json：每份样例的期望字段值，供评测（app/eval/run_eval.py）与人工核对使用。

用法：python data/gen_samples.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.config import settings  # noqa: E402

SAMPLES = [
    {"id": "sample01", "record_no": "20260001", "name": "张伟", "gender": "男", "birth_date": "1985-03-12",
     "admission_date": "2026-05-10", "discharge_date": "2026-05-18", "admission_dept": "心血管内科",
     "discharge_dept": "心血管内科", "main_diagnosis": "原发性高血压", "main_icd": "I10",
     "secondary_diagnosis": "2型糖尿病", "surgery": "冠状动脉造影术", "attending_doctor": "李梅",
     "hospital_days": "9", "payment_type": "医保"},
    {"id": "sample02", "record_no": "20260002", "name": "王芳", "gender": "女", "birth_date": "1990-07-25",
     "admission_date": "2026-06-02", "discharge_date": "2026-06-09", "admission_dept": "呼吸内科",
     "discharge_dept": "呼吸内科", "main_diagnosis": "肺炎", "main_icd": "J18.9",
     "secondary_diagnosis": "缺铁性贫血", "surgery": "", "attending_doctor": "陈刚",
     "hospital_days": "8", "payment_type": "医保"},
    {"id": "sample03", "record_no": "20260003", "name": "刘强", "gender": "男", "birth_date": "1972-11-03",
     "admission_date": "2026-06-15", "discharge_date": "2026-06-25", "admission_dept": "神经内科",
     "discharge_dept": "康复医学科", "main_diagnosis": "脑梗死", "main_icd": "I63.9",
     "secondary_diagnosis": "高脂血症", "surgery": "脑血管介入取栓术", "attending_doctor": "赵敏",
     "hospital_days": "11", "payment_type": "医保"},
    {"id": "sample04", "record_no": "20260004", "name": "陈静", "gender": "女", "birth_date": "1998-01-20",
     "admission_date": "2026-07-01", "discharge_date": "2026-07-05", "admission_dept": "普外科",
     "discharge_dept": "普外科", "main_diagnosis": "急性阑尾炎", "main_icd": "K35.8",
     "secondary_diagnosis": "", "surgery": "腹腔镜阑尾切除术", "attending_doctor": "孙涛",
     "hospital_days": "5", "payment_type": "自费"},
    {"id": "sample05", "record_no": "20260005", "name": "李娜", "gender": "女", "birth_date": "1968-09-15",
     "admission_date": "2026-07-08", "discharge_date": "2026-07-12", "admission_dept": "骨科",
     "discharge_dept": "骨科", "main_diagnosis": "膝关节病", "main_icd": "M17.9",
     "secondary_diagnosis": "骨质疏松", "surgery": "膝关节置换术", "attending_doctor": "周杰",
     "hospital_days": "5", "payment_type": "医保"},
    {"id": "sample06", "record_no": "20260006", "name": "赵磊", "gender": "男", "birth_date": "1955-04-08",
     "admission_date": "2026-07-20", "discharge_date": "2026-07-28", "admission_dept": "肿瘤科",
     "discharge_dept": "肿瘤科", "main_diagnosis": "肺恶性肿瘤", "main_icd": "C34.9",
     "secondary_diagnosis": "慢性阻塞性肺病", "surgery": "胸腔镜下肺叶切除术", "attending_doctor": "吴艳",
     "hospital_days": "9", "payment_type": "医保"},
    {"id": "sample07", "record_no": "20260007", "name": "孙丽", "gender": "女", "birth_date": "2001-12-30",
     "admission_date": "2026-08-01", "discharge_date": "2026-08-03", "admission_dept": "消化内科",
     "discharge_dept": "消化内科", "main_diagnosis": "慢性胃炎", "main_icd": "K29.5",
     "secondary_diagnosis": "胃食管反流病", "surgery": "胃镜检查", "attending_doctor": "郑伟",
     "hospital_days": "3", "payment_type": "医保"},
    {"id": "sample08", "record_no": "20260008", "name": "周敏", "gender": "女", "birth_date": "1982-06-18",
     "admission_date": "2026-08-05", "discharge_date": "2026-08-07", "admission_dept": "泌尿外科",
     "discharge_dept": "泌尿外科", "main_diagnosis": "肾结石", "main_icd": "N20.0",
     "secondary_diagnosis": "泌尿道感染", "surgery": "体外冲击波碎石术", "attending_doctor": "王磊",
     "hospital_days": "3", "payment_type": "医保"},
    {"id": "sample09", "record_no": "20260009", "name": "吴刚", "gender": "男", "birth_date": "1960-02-11",
     "admission_date": "2026-08-10", "discharge_date": "2026-08-20", "admission_dept": "心血管内科",
     "discharge_dept": "心血管内科", "main_diagnosis": "心房颤动", "main_icd": "I48",
     "secondary_diagnosis": "心力衰竭", "surgery": "心脏射频消融术", "attending_doctor": "李梅",
     "hospital_days": "11", "payment_type": "公费"},
    {"id": "sample10", "record_no": "20260010", "name": "郑浩", "gender": "男", "birth_date": "1995-10-05",
     "admission_date": "2026-08-15", "discharge_date": "2026-08-18", "admission_dept": "内分泌科",
     "discharge_dept": "内分泌科", "main_diagnosis": "2型糖尿病", "main_icd": "E11.9",
     "secondary_diagnosis": "高脂血症", "surgery": "", "attending_doctor": "冯雪",
     "hospital_days": "4", "payment_type": "医保"},
    {"id": "sample11", "record_no": "20260011", "name": "冯雪", "gender": "女", "birth_date": "1978-08-22",
     "admission_date": "2026-08-20", "discharge_date": "2026-08-26", "admission_dept": "骨科",
     "discharge_dept": "骨科", "main_diagnosis": "腰背痛", "main_icd": "M54.5",
     "secondary_diagnosis": "腰椎间盘突出", "surgery": "腰椎减压术", "attending_doctor": "周杰",
     "hospital_days": "7", "payment_type": "商保"},
    {"id": "sample12", "record_no": "20260012", "name": "刘洋", "gender": "男", "birth_date": "2015-05-30",
     "admission_date": "2026-08-22", "discharge_date": "2026-08-25", "admission_dept": "儿科",
     "discharge_dept": "儿科", "main_diagnosis": "肺炎", "main_icd": "J18.9",
     "secondary_diagnosis": "哮喘", "surgery": "", "attending_doctor": "杨帆",
     "hospital_days": "4", "payment_type": "医保"},
]

FIELD_ORDER = ["record_no", "name", "gender", "birth_date", "admission_date", "discharge_date",
               "admission_dept", "discharge_dept", "main_diagnosis", "main_icd",
               "secondary_diagnosis", "surgery", "attending_doctor", "hospital_days", "payment_type"]

LABELS = {
    "record_no": "病案号", "name": "姓名", "gender": "性别", "birth_date": "出生日期",
    "admission_date": "入院日期", "discharge_date": "出院日期", "admission_dept": "入院科室",
    "discharge_dept": "出院科室", "main_diagnosis": "主要诊断", "main_icd": "主要诊断ICD编码",
    "secondary_diagnosis": "次要诊断", "surgery": "手术操作", "attending_doctor": "主治医师",
    "hospital_days": "住院天数", "payment_type": "费用类别",
}


def _kv_lines(s: dict) -> list[str]:
    lines = []
    for k in FIELD_ORDER:
        lines.append(f"{LABELS[k]}: {s[k]}")
    return lines


def _render_txt(s: dict) -> str:
    return "\n".join(_kv_lines(s)) + "\n"


def _render_md(s: dict) -> str:
    lines = [f"# 病案首页（样例 {s['id']}）", ""]
    lines += [f"- {LABELS[k]}：{s[k]}" for k in FIELD_ORDER]
    return "\n".join(lines) + "\n"


def _render_pdf(s: dict, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    c.setFont("STSong-Light", 14)
    c.drawString(60, height - 60, f"病案首页（样例 {s['id']}）")
    c.setFont("STSong-Light", 11)
    y = height - 100
    for line in _kv_lines(s):
        c.drawString(60, y, line)
        y -= 22
    c.save()


def _render_docx(s: dict, path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(f"病案首页（样例 {s['id']}）", level=1)
    table = doc.add_table(rows=len(FIELD_ORDER), cols=2)
    table.style = "Table Grid"
    for i, k in enumerate(FIELD_ORDER):
        table.rows[i].cells[0].text = LABELS[k]
        table.rows[i].cells[1].text = s[k]
    doc.save(str(path))


def _render_xlsx(s: dict, path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "病案首页"
    ws.append(["字段", "值"])
    for k in FIELD_ORDER:
        ws.append([LABELS[k], s[k]])
    wb.save(str(path))


def _render_png(s: dict, path: Path) -> bool:
    """用 PIL 画一张表单样式图片（模拟扫描件）。无中文字体时跳过。"""
    from PIL import Image, ImageDraw, ImageFont

    font_path = None
    for p in (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf",
              r"C:\Windows\Fonts\simsun.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"):
        if Path(p).exists():
            font_path = p
            break
    if not font_path:
        return False
    font_title = ImageFont.truetype(font_path, 28)
    font_body = ImageFont.truetype(font_path, 18)
    lines = [f"病案首页（样例 {s['id']}）"] + _kv_lines(s)
    w, h = 640, 90 + len(lines) * 34
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    # 表单横线
    draw.rectangle([20, 20, w - 20, h - 20], outline="black", width=2)
    y = 55
    draw.text((40, 15), lines[0], font=font_title, fill="black")
    for line in lines[1:]:
        draw.line([40, y, w - 40, y], fill="black", width=1)
        draw.text((50, y - 24), line, font=font_body, fill="black")
        y += 34
    img.save(str(path))
    return True


def generate_samples(out_dir: Path | None = None) -> dict:
    out_dir = out_dir or settings.samples_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    expected = {}
    for s in SAMPLES:
        sid = s["id"]
        # txt / md
        (out_dir / f"{sid}.txt").write_text(_render_txt(s), encoding="utf-8")
        (out_dir / f"{sid}.md").write_text(_render_md(s), encoding="utf-8")
        # pdf
        _render_pdf(s, out_dir / f"{sid}.pdf")
        # docx
        _render_docx(s, out_dir / f"{sid}.docx")
        # xlsx
        _render_xlsx(s, out_dir / f"{sid}.xlsx")
        # png（尽力而为）
        _render_png(s, out_dir / f"{sid}.png")
        expected[sid] = {k: s[k] for k in FIELD_ORDER}
    (out_dir / "expected.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2), encoding="utf-8")
    return expected


if __name__ == "__main__":
    exp = generate_samples()
    print(f"已生成 12 份样例 × 6 格式 → {settings.samples_dir}")
    print(f"期望标注：{len(exp)} 份（expected.json）")
